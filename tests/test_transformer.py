from io import BytesIO

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from src.transformer import TransformationError, transform_distributor_report


def _destination_workbook_bytes(include_preview: bool = False) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = "Existing content"
    summary["B2"] = "=1+1"
    summary["A1"].font = Font(bold=True)
    summary.column_dimensions["A"].width = 24
    summary.merge_cells("C3:D3")
    summary["C3"] = "Merged heading"
    workbook.create_sheet("September 2026")
    if include_preview:
        preview = workbook.create_sheet("Automation Preview")
        preview["A1"] = "Old preview"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_transformer_preserves_original_and_adds_preview() -> None:
    source_dataframe = pd.DataFrame(
        {
            "Distributor": ["Acme", "Acme"],
            "Sales": [100, 250],
        }
    )

    result = transform_distributor_report(
        source_df=source_dataframe,
        selected_distributor="Acme",
        destination_workbook_bytes=_destination_workbook_bytes(),
        destination_sheet_name="September 2026",
    )

    workbook = load_workbook(BytesIO(result), data_only=False)
    summary = workbook["Summary"]
    preview = workbook["Automation Preview"]

    assert summary["A1"].value == "Existing content"
    assert summary["A1"].font.bold is True
    assert summary["B2"].value == "=1+1"
    assert summary.column_dimensions["A"].width == 24
    assert "C3:D3" in summary.merged_cells
    assert summary["C3"].value == "Merged heading"
    assert preview["A1"].value == "Automation Setup Complete"
    assert preview["A3"].value == "Distributor: Acme"
    assert preview["A4"].value == "Destination Sheet: September 2026"
    assert preview["A5"].value == "Rows Found: 2"
    assert preview["A7"].value == (
        "Actual mapping logic has not been configured yet."
    )
    assert preview["A9"].value == "Distributor"
    assert preview["B10"].value == 100
    assert preview["B11"].value == 250
    workbook.close()


def test_transformer_replaces_existing_preview_sheet() -> None:
    result = transform_distributor_report(
        source_df=pd.DataFrame({"Distributor": ["Acme"]}),
        selected_distributor="Acme",
        destination_workbook_bytes=_destination_workbook_bytes(include_preview=True),
        destination_sheet_name="September 2026",
    )

    workbook = load_workbook(BytesIO(result))
    assert workbook.sheetnames.count("Automation Preview") == 1
    assert workbook["Automation Preview"]["A1"].value == "Automation Setup Complete"
    workbook.close()


def test_transformer_rejects_missing_destination_sheet() -> None:
    with pytest.raises(TransformationError, match="could not be found"):
        transform_distributor_report(
            source_df=pd.DataFrame({"Distributor": ["Acme"]}),
            selected_distributor="Acme",
            destination_workbook_bytes=_destination_workbook_bytes(),
            destination_sheet_name="Missing",
        )
