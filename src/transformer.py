from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.excel_writer import (
    save_workbook_to_bytes,
    write_automation_preview_sheet,
)
from src.validators import ValidationError, validate_source_dataframe


class TransformationError(ValueError):
    """Raised when the destination workbook cannot be transformed."""


def transform_distributor_report(
    source_df: pd.DataFrame,
    selected_distributor: str,
    destination_workbook_bytes: bytes,
    destination_sheet_name: str,
) -> bytes:
    """Generate a destination workbook using the current placeholder flow."""
    try:
        validate_source_dataframe(source_df)
    except ValidationError as exc:
        raise TransformationError(str(exc)) from exc

    if not selected_distributor.strip():
        raise TransformationError(
            "The selected distributor is not available in the source data."
        )

    try:
        workbook = load_workbook(
            BytesIO(destination_workbook_bytes),
            data_only=False,
            keep_links=True,
        )
    except Exception as exc:
        raise TransformationError(
            "Unable to open the destination workbook."
        ) from exc

    try:
        if destination_sheet_name not in workbook.sheetnames:
            raise TransformationError(
                "The selected destination sheet could not be found."
            )

        # TODO:
        # Replace this placeholder with the actual distributor
        # calculation and destination-sheet mapping rules.
        write_automation_preview_sheet(
            workbook=workbook,
            source_df=source_df,
            selected_distributor=selected_distributor,
            destination_sheet_name=destination_sheet_name,
        )
        return save_workbook_to_bytes(workbook)
    except TransformationError:
        raise
    except Exception as exc:
        raise TransformationError(
            "Unable to generate the updated workbook."
        ) from exc
    finally:
        workbook.close()

