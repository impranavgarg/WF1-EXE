from collections.abc import Collection, Sequence
from io import BytesIO
from typing import Hashable

import pandas as pd
from openpyxl import load_workbook


class ValidationError(ValueError):
    """Raised when user-provided report inputs are not valid."""


def validate_source_dataframe(source_df: pd.DataFrame) -> None:
    """Ensure the selected source sheet contains tabular data."""
    if not isinstance(source_df, pd.DataFrame) or source_df.empty:
        raise ValidationError("No data was found in the selected sheet.")
    if len(source_df.columns) == 0:
        raise ValidationError("No data was found in the selected sheet.")


def validate_distributor_column(
    source_df: pd.DataFrame, distributor_column: Hashable | None
) -> None:
    """Ensure the distributor column is present in the source data."""
    if distributor_column is None or distributor_column not in source_df.columns:
        raise ValidationError(
            "The selected distributor column is not available in the source data."
        )


def validate_selected_distributor(
    selected_distributor: str | None,
    available_distributors: Collection[str],
) -> None:
    """Ensure the selected distributor came from the source workbook."""
    if (
        selected_distributor is None
        or selected_distributor not in available_distributors
    ):
        raise ValidationError(
            "The selected distributor is not available in the source data."
        )


def validate_destination_workbook(destination_workbook_bytes: bytes) -> None:
    """Ensure the uploaded destination bytes contain an openable workbook."""
    if not destination_workbook_bytes:
        raise ValidationError("Unable to open the destination workbook.")

    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(destination_workbook_bytes),
            read_only=True,
            data_only=False,
        )
        if not workbook.sheetnames:
            raise ValidationError("Unable to open the destination workbook.")
    except ValidationError:
        raise
    except Exception as exc:
        # openpyxl can surface several ZIP/XML exception types for corrupt files.
        raise ValidationError("Unable to open the destination workbook.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def validate_destination_sheet(
    available_sheets: Sequence[str], destination_sheet: str | None
) -> None:
    """Ensure the selected destination sheet exists in the workbook."""
    if destination_sheet is None or destination_sheet not in available_sheets:
        raise ValidationError(
            "The selected destination sheet could not be found."
        )

