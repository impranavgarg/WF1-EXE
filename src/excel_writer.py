import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


PREVIEW_SHEET_NAME = "Automation Preview"
INVALID_FILENAME_CHARACTERS = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


def build_output_filename(
    distributor_name: str, generated_on: date | None = None
) -> str:
    """Build a filesystem-safe output filename for the generated report."""
    safe_name = re.sub(r"\s+", "_", distributor_name.strip())
    safe_name = INVALID_FILENAME_CHARACTERS.sub("_", safe_name)
    safe_name = re.sub(r"_+", "_", safe_name).strip(" ._")
    safe_name = safe_name[:100].rstrip(" ._") or "Distributor"
    report_date = generated_on or date.today()
    return f"{safe_name}_Updated_{report_date.isoformat()}.xlsx"


def write_automation_preview_sheet(
    workbook: Workbook,
    source_df: pd.DataFrame,
    selected_distributor: str,
    destination_sheet_name: str,
) -> Worksheet:
    """Replace the placeholder preview sheet and populate it with source data."""
    preview_index = len(workbook.sheetnames)
    if PREVIEW_SHEET_NAME in workbook.sheetnames:
        preview_index = workbook.sheetnames.index(PREVIEW_SHEET_NAME)
        workbook.remove(workbook[PREVIEW_SHEET_NAME])

    worksheet = workbook.create_sheet(PREVIEW_SHEET_NAME, preview_index)
    worksheet["A1"] = "Automation Setup Complete"
    worksheet["A1"].font = Font(size=14, bold=True)
    worksheet["A3"] = f"Distributor: {selected_distributor}"
    worksheet["A4"] = f"Destination Sheet: {destination_sheet_name}"
    worksheet["A5"] = f"Rows Found: {len(source_df)}"
    worksheet["A7"] = "Actual mapping logic has not been configured yet."
    worksheet["A7"].font = Font(italic=True)

    _write_dataframe(worksheet, source_df, start_row=9)
    return worksheet


def save_workbook_to_bytes(workbook: Workbook) -> bytes:
    """Serialize a workbook to in-memory XLSX bytes."""
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_dataframe(
    worksheet: Worksheet, dataframe: pd.DataFrame, start_row: int
) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for column_index, column_name in enumerate(dataframe.columns, start=1):
        cell = worksheet.cell(
            row=start_row,
            column=column_index,
            value=str(column_name),
        )
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row_index, row_values in enumerate(
        dataframe.itertuples(index=False, name=None), start=start_row + 1
    ):
        for column_index, value in enumerate(row_values, start=1):
            worksheet.cell(
                row=row_index,
                column=column_index,
                value=_excel_cell_value(value),
            )

    if len(dataframe.columns) > 0:
        worksheet.auto_filter.ref = (
            f"A{start_row}:"
            f"{get_column_letter(len(dataframe.columns))}{start_row + len(dataframe)}"
        )
        worksheet.freeze_panes = f"A{start_row + 1}"
        _set_preview_column_widths(worksheet, dataframe)


def _excel_cell_value(value: Any) -> Any:
    try:
        if bool(pd.isna(value)):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, pd.Timedelta):
        return str(value)
    if hasattr(value, "item") and callable(value.item):
        return value.item()
    if isinstance(value, (date, datetime, str, int, float, bool)):
        return value
    return str(value)


def _set_preview_column_widths(
    worksheet: Worksheet, dataframe: pd.DataFrame
) -> None:
    for column_index, column_name in enumerate(dataframe.columns, start=1):
        values = dataframe.iloc[:100, column_index - 1]
        content_lengths = [
            len(str(value)) for value in values if _excel_cell_value(value) is not None
        ]
        longest_value = max([len(str(column_name)), *content_lengths])
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(longest_value + 2, 12), 40
        )

