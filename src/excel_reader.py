import re
from collections.abc import Hashable
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook


class ExcelReadError(ValueError):
    """Raised when an uploaded Excel workbook cannot be read."""


def read_sheet_names(workbook_bytes: bytes) -> list[str]:
    """Return workbook sheet names in their existing order."""
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(workbook_bytes),
            read_only=True,
            data_only=False,
        )
        return list(workbook.sheetnames)
    except Exception as exc:
        raise ExcelReadError("Unable to read this Excel file.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def read_source_sheet(workbook_bytes: bytes, sheet_name: str) -> pd.DataFrame:
    """Read one source workbook sheet into a dataframe."""
    try:
        return pd.read_excel(
            BytesIO(workbook_bytes),
            sheet_name=sheet_name,
            engine="openpyxl",
        )
    except Exception as exc:
        raise ExcelReadError("Unable to read this Excel file.") from exc


def clean_distributor_name(value: object) -> str | None:
    """Convert a source value to a trimmed distributor name, or None."""
    if value is None:
        return None

    try:
        if bool(pd.isna(value)):
            return None
    except (TypeError, ValueError):
        # Source dataframe cells are expected to be scalar values. If a custom
        # scalar cannot be checked by pandas, its string form is still usable.
        pass

    cleaned_value = str(value).strip()
    return cleaned_value or None


def extract_unique_distributor_names(
    source_df: pd.DataFrame, distributor_column: Hashable
) -> list[str]:
    """Return cleaned, unique distributor names in alphabetical order."""
    if distributor_column not in source_df.columns:
        raise KeyError(f"Unknown distributor column: {distributor_column}")

    names = {
        cleaned_name
        for value in source_df[distributor_column]
        if (cleaned_name := clean_distributor_name(value)) is not None
    }
    return sorted(names, key=lambda name: (name.casefold(), name))


def filter_distributor_rows(
    source_df: pd.DataFrame,
    distributor_column: Hashable,
    selected_distributor: str,
) -> pd.DataFrame:
    """Return source rows whose cleaned value matches the selected name."""
    selected_name = clean_distributor_name(selected_distributor)
    if selected_name is None:
        return source_df.iloc[0:0].copy()

    cleaned_values = source_df[distributor_column].map(clean_distributor_name)
    return source_df.loc[cleaned_values == selected_name].copy()


def suggest_current_month_sheet(
    sheet_names: list[str], current_date: date | None = None
) -> str | None:
    """Choose the best current-month sheet, falling back to the first sheet."""
    if not sheet_names:
        return None

    today = current_date or date.today()
    full_month = today.strftime("%B").casefold()
    short_month = today.strftime("%b").casefold()
    year = str(today.year)

    def normalize(value: str) -> str:
        return re.sub(r"[\s_-]+", " ", value.strip().casefold())

    exact_scores = {
        f"{full_month} {year}": 100,
        f"{short_month} {year}": 95,
        full_month: 90,
        short_month: 85,
    }

    best_sheet = sheet_names[0]
    best_score = 0
    for sheet_name in sheet_names:
        normalized_name = normalize(sheet_name)
        score = exact_scores.get(normalized_name, 0)
        tokens = set(normalized_name.split())
        if score == 0 and full_month in tokens and year in tokens:
            score = 80
        elif score == 0 and short_month in tokens and year in tokens:
            score = 75
        elif score == 0 and full_month in tokens:
            score = 70
        elif score == 0 and short_month in tokens:
            score = 65

        if score > best_score:
            best_sheet = sheet_name
            best_score = score

    return best_sheet

